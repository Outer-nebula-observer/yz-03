# -*- coding: utf-8 -*-
"""
pptx2svg.py —— SVG -> PPTX 的反向转换（用于把 PPTX 里手工改过的版本同步回 SVG/PNG）

用法：
    python pptx2svg.py fig1_architecture.pptx [输出.svg]

为什么需要它：在 PowerPoint 里改完再让脚本重新生成 SVG，会把改动覆盖掉。
这里反过来，以 PPTX 为唯一真相，精确反推出 SVG，三个文件（svg/png/pptx）因此严格一致。

换算与 svg2pptx.py 严格互逆：
    1px = 9525 EMU（96dpi）;  pt -> px: ÷0.75
    非旋转文本:  SVG 基线 y = PPTX 顶边 top + ASCENT_EM × size_px
    旋转文本:    SVG 中心 y = PPTX 顶边 top + size_px × 0.8
    圆角矩形:    rx = adjustment × min(w, h)
"""

import sys
import re
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

PX = 9525
ASCENT_EM = 1.079


def px(v):
    return round(float(v) / PX, 2)


def rgb_of(colorformat):
    try:
        return str(colorformat.rgb)
    except Exception:
        return None


def shape_fill(sh):
    try:
        if sh.fill.type is None:
            return None
        return rgb_of(sh.fill.fore_color)
    except Exception:
        return None


def shape_line(sh):
    """返回 (hex, 磅值) 或 (None, None)。"""
    try:
        ln = sh.line
        if ln.fill.type is None:
            return None, None
        hexv = rgb_of(ln.color)
        w = ln.width.pt if ln.width else 0.75
        return hexv, round(w, 2)
    except Exception:
        return None, None


def has_arrow(sh):
    ln = sh._element.spPr.find(qn("a:ln"))
    return ln is not None and ln.find(qn("a:tailEnd")) is not None


def freeform_points(sh):
    """从 custGeom 里取出折线顶点（已是画布 px 坐标）。"""
    spPr = sh._element.spPr
    cg = spPr.find(qn("a:custGeom"))
    if cg is None:
        return []
    p = cg.find(qn("a:pathLst")).find(qn("a:path"))
    pw, ph = int(p.get("w") or 1), int(p.get("h") or 1)
    L, T = sh.left / PX, sh.top / PX
    W, H = sh.width / PX, sh.height / PX
    out = []
    for child in p:
        tag = child.tag.split("}")[1]
        if tag in ("moveTo", "lnTo"):
            pt = child.find(qn("a:pt"))
            out.append((round(L + int(pt.get("x")) / pw * W, 1),
                        round(T + int(pt.get("y")) / ph * H, 1)))
    return out


def text_info(sh):
    """返回 (文本, size_px, color, weight, align) 或 None。"""
    if not sh.has_text_frame:
        return None
    s = sh.text_frame.text
    if not s.strip():
        return None
    para = sh.text_frame.paragraphs[0]
    if not para.runs:
        return None
    r = para.runs[0]
    size_pt = r.font.size.pt if r.font.size else 18.0
    weight = 600 if (r.font.name and "Semibold" in r.font.name) else 400
    align = {PP_ALIGN.CENTER: "middle", PP_ALIGN.RIGHT: "end"}.get(para.alignment, "start")
    return s, round(size_pt / 0.75, 2), rgb_of(r.font.color) or "0B0B0B", weight, align


def convert(pptx_path, svg_path):
    prs = Presentation(pptx_path)
    slide = prs.slides[0]
    W, H = px(prs.slide_width), px(prs.slide_height)

    layers = {"bg": [], "shape": [], "arrow": [], "text": []}
    markers = set()

    for sh in slide.shapes:
        st = str(sh.shape_type)
        L, T = px(sh.left), px(sh.top)
        w, h = px(sh.width), px(sh.height)
        fill = shape_fill(sh)
        stroke, lw = shape_line(sh)

        if "AUTO_SHAPE" in st and w >= W and h >= H:
            layers["bg"].append(
                '<rect x="%.2f" y="%.2f" width="%.2f" height="%.2f" fill="#%s"/>'
                % (L, T, w, h, fill or "FFFFFF"))

        elif "AUTO_SHAPE" in st:
            try:
                adj = sh.adjustments[0]
            except Exception:
                adj = 0.0
            rx = round(adj * min(w, h), 2)
            a = ' rx="%.2f"' % rx if rx > 0.01 else ""
            f = ' fill="#%s"' % fill if fill else ' fill="none"'
            s_ = ' stroke="#%s" stroke-width="%.2f"' % (stroke, lw / 0.75) if stroke else ""
            layers["shape"].append(
                '<rect x="%.2f" y="%.2f" width="%.2f" height="%.2f"%s%s%s/>'
                % (L, T, w, h, a, f, s_))

        elif "FREEFORM" in st or "LINE" in st:
            pts = freeform_points(sh) if "FREEFORM" in st else None
            if not pts:
                if "LINE" in st:
                    # 连接线：left/top/width/height + flip 决定端点
                    xfrm = sh._element.spPr.find(qn("a:xfrm"))
                    fh = xfrm is not None and xfrm.get("flipH") == "1"
                    fv = xfrm is not None and xfrm.get("flipV") == "1"
                    x1, y1 = (L + w, T) if fh else (L, T)
                    x2, y2 = (L, T + h) if fh else (L + w, T + h)
                    if fv:
                        y1, y2 = y2, y1
                    pts = [(round(x1, 1), round(y1, 1)), (round(x2, 1), round(y2, 1))]
                else:
                    continue
            d = "M" + " L".join("%.1f,%.1f" % p for p in pts)
            stroke = stroke or "52514E"
            # marker 必须按线色分开定义，否则蓝线的箭头会被染成灰色
            arrow = ' marker-end="url(#arw-%s)"' % stroke if has_arrow(sh) else ""
            if has_arrow(sh):
                markers.add(stroke)
            layers["arrow"].append(
                '<path d="%s" fill="none" stroke="#%s" stroke-width="%.2f"%s/>'
                % (d, stroke, (lw or 1.5) / 0.75, arrow))

        else:
            info = text_info(sh)
            if not info:
                continue
            s, size, color, weight, align = info
            rot = int(sh.rotation)
            if rot:
                cx, cy = L + w / 2.0, T + size * 0.8
                head = ' transform="translate(%.2f,%.2f) rotate(%d)"' % (cx, cy, rot)
                anchor = "middle"
            else:
                cy = T + ASCENT_EM * size
                head = ""
                anchor = align
                if anchor == "middle":
                    cx = L + w / 2.0
                elif anchor == "end":
                    cx = L + w
                else:
                    cx = L
            attrs = ('x="%.2f" y="%.2f" font-size="%.2fpx" font-weight="%d" fill="#%s"'
                     % (cx, cy, size, weight, color))
            if anchor != "start":
                attrs += ' text-anchor="%s"' % anchor
            layers["text"].append("<text %s%s>%s</text>" % (attrs, head, escape(s)))

    out = ['<svg xmlns="http://www.w3.org/2000/svg" width="%.0f" height="%.0f" '
           'viewBox="0 0 %.0f %.0f">' % (W, H, W, H), '  <defs>']
    for c in sorted(markers):
        out.append('    <marker id="arw-%s" viewBox="0 0 10 10" refX="9" refY="5" '
                   'markerWidth="5" markerHeight="5" orient="auto-start-reverse">' % c)
        out.append('      <path d="M0,0 L10,5 L0,10 z" fill="#%s"/>' % c)
        out.append('    </marker>')
    out.append('  </defs>')
    for key in ("bg", "shape", "arrow", "text"):
        for line in layers[key]:
            out.append("  " + line)
    out.append("</svg>")

    with open(svg_path, "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")
    print("  %-28s -> %-28s %d 底色 / %d 形状 / %d 箭头 / %d 文字"
          % (pptx_path, svg_path, len(layers["bg"]), len(layers["shape"]),
             len(layers["arrow"]), len(layers["text"])))


if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else src.rsplit(".", 1)[0] + ".svg"
    convert(src, dst)
